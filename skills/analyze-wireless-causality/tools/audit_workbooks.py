"""Stream XLSX tables; profile values without interpreting names as measurements."""
import argparse
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from common import number, sha256, table_rows, write_json, ensure_output_safe


def aggregate_status(items):
    """Retain failures while allowing successful sibling sheets/files to finish."""
    statuses = [item['status'] for item in items]
    if all(status == 'completed' for status in statuses):
        return 'completed'
    if all(status == 'failed' for status in statuses):
        return 'failed'
    return 'partial'


def audit(path, max_rows=None, header_row=1, sheet=None):
    if max_rows is not None and (type(max_rows) is not int or max_rows < 1):
        raise ValueError("max_rows must be a positive integer")
    if type(header_row) is not int or header_row < 1:
        raise ValueError("header_row must be a positive integer")
    import openpyxl
    before = sha256(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [(ws.title, ws.max_row, ws.max_column) for ws in workbook if sheet is None or ws.title == sheet]
    finally:
        workbook.close()
    if not sheets:
        raise ValueError('Selected sheet is absent')
    result = {'path':str(Path(path).resolve()), 'sha256':before, 'sheets':[],
              'mode':'full' if max_rows is None else 'sample', 'header_row':header_row,
              'max_rows':max_rows, 'sheet_selection':sheet,
              'scan_scope':'all_sheets' if sheet is None else 'selected_sheet',
              'formula_policy':'cached values only; formula caches may be absent or stale'}
    for sheet, declared_rows, declared_cols in sheets:
        profile, count, issue, issue_type = {}, 0, None, None
        last_source_row, boundary = None, 'end_of_sheet'
        rows = table_rows(path, sheet, header_row)
        try:
            # Check the requested limit before next(): another row can be corrupt.
            while max_rows is None or count < max_rows:
                try:
                    rowno, row = next(rows)
                except StopIteration:
                    break
                except Exception as exc:
                    # Only reader failures become diagnostics. Input, integrity and
                    # output-safety validation remain outside this recovery boundary.
                    issue, issue_type, boundary = str(exc), type(exc).__name__, 'read_error'
                    break
                count += 1
                last_source_row = rowno
                for name, value in row.items():
                    col = profile.setdefault(name, {'present':0,'numeric':0,'min':None,'max':None,
                                                     'examples':[], 'bounded_values':Counter()})
                    if value is None or value == '':
                        continue
                    col['present'] += 1
                    text = str(value)
                    if len(col['examples']) < 4 and text[:200] not in col['examples']:
                        col['examples'].append(text[:200])
                    if text in col['bounded_values'] or len(col['bounded_values']) < 64:
                        col['bounded_values'][text] += 1
                    n = number(value)
                    if n is not None:
                        col['numeric'] += 1
                        col['min'] = n if col['min'] is None else min(n, col['min'])
                        col['max'] = n if col['max'] is None else max(n, col['max'])
            else:
                boundary = 'max_rows'
        finally:
            rows.close()
        for col in profile.values():
            col['missing'] = count - col['present']
            col['first_64_values_counts'] = dict(col.pop('bounded_values'))
        result['sheets'].append({'sheet':sheet,'declared_rows':declared_rows,
                                'declared_columns':declared_cols,'rows_scanned':count,
                                'status':('partial' if count else 'failed') if issue_type else 'completed',
                                'issue':issue,'issue_type':issue_type,'columns':profile,
                                'scan':{'max_rows':max_rows,'limit_unit':'nonempty_data_rows',
                                        'boundary':boundary,'last_source_row':last_source_row,
                                        'full_sheet_scanned':boundary == 'end_of_sheet'}})
    if sha256(path) != before:
        raise RuntimeError('Source changed during audit')
    result['status'] = aggregate_status(result['sheets'])
    result['full_scan_completed'] = all(s['scan']['full_sheet_scanned'] for s in result['sheets'])
    return result


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--input', nargs='+', required=True)
    p.add_argument('--output', required=True)
    p.add_argument('--max-rows', type=int)
    p.add_argument('--header-row', type=int, default=1)
    p.add_argument('--sheet')
    a=p.parse_args()
    if a.max_rows is not None and a.max_rows < 1:
        p.error('--max-rows must be positive')
    if Path(a.output).suffix.lower() != '.json':
        p.error('Output must be a JSON file')
    if any(Path(x).suffix.lower() != '.xlsx' for x in a.input):
        p.error('This audit reads .xlsx only')
    if any(Path(x).resolve() == Path(a.output).resolve() for x in a.input):
        p.error('Output cannot overwrite a source')
    ensure_output_safe(a.output,a.input,'.json')
    results=[]
    for path in a.input:
        r=audit(path,a.max_rows,a.header_row,a.sheet); results.append(r)
        print(Path(path).name, r['status'], [(s['sheet'],s['rows_scanned']) for s in r['sheets']], flush=True)
    status = aggregate_status(results)
    write_json(a.output, {'created_utc':datetime.now(timezone.utc).isoformat(),
                         'status':status,'mode':'full' if a.max_rows is None else 'sample',
                         'max_rows':a.max_rows,'header_row':a.header_row,'sheet_selection':a.sheet,
                         'scan_scope':'all_sheets' if a.sheet is None else 'selected_sheet',
                         'full_scan_completed':all(r['full_scan_completed'] for r in results),
                         'workbooks':results},a.input)
    return 0 if status == 'completed' else 1

if __name__ == '__main__':
    raise SystemExit(main())
