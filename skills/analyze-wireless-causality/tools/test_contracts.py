"""Adversarial synthetic cases for boundaries, missingness and evidence integrity."""
import copy
import json
import math
import os
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl
from audit_workbooks import audit
from build_windows import build
from common import ensure_output_safe, read_json, sha256, table_rows, validate_windows
from compare_runs import paired_compare, summarize
from detect_events import detect
from plot_windows import ecdf_points, plot
from wifi_frames import decode_compressed_ba, observed_retry_chains
from write_report import render, validate_evidence, verify_locator


class ContractTests(unittest.TestCase):
    def setUp(self):
        self.base = Path(__file__).resolve().parents[1] / 'logs' / 'tool_tests'
        self.base.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=self.base, prefix='contract_')
        self.root = Path(self.temp.name)

    def tearDown(self):
        if self.root.resolve().parent != self.base.resolve():
            raise RuntimeError('Unexpected cleanup path')
        self.temp.cleanup()

    def table(self, rows, headers=('t', 'end', 'v', 'fail', 'trials', 'kind')):
        path = self.root / 'source.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Arbitrary'
        ws.append(list(headers))
        for row in rows:
            ws.append(row)
        wb.save(path)
        wb.close()
        return path

    def request(self, rows, event=False, kind='interval_amount', width=1, complete=True):
        path = self.table(rows)
        timing = {'start': 't', 'scale': 1, 'origin': 0, 'split_policy': 'reject'}
        if event:
            timing['observation_interval'] = [0, 1]
        else:
            timing['end'] = 'end'
        return {'schema_version': 1, 'window_s': width, 'sources': [{
            'path': str(path), 'sheet': 'Arbitrary', 'capture_complete': complete,
            'metadata': {'run_id': 'r', 'entity': 'A', 'direction': 'outbound', 'condition': 'c', 'pair_id': 'p'},
            'time': timing, 'metrics': [{'name': 'metric', 'unit': 'u', 'kind': kind, 'column': 'v',
                                      'numerator': 'fail', 'denominator': 'trials', 'scale': 1,
                                      'definition': 'Declared synthetic measurement'}]}]}

    def windows(self, values):
        return [{'run_id': 'r', 'entity': 'A', 'direction': 'outbound', 'condition': 'c', 'pair_id': 'p',
                 'metric': 'm', 'unit': 'u', 'start_s': i, 'end_s': i + 1, 'value': v, 'coverage': 1,
                 'kind': 'interval_mean', 'definition': 'Observed duration-weighted outcome', 'source_rows': [i + 2]}
                for i, v in enumerate(values)]

    def report(self):
        evidence = self.root / 'evidence.json'
        evidence.write_text('{"rows":[{"value":1}]}', encoding='utf-8')
        figure = self.root / 'plot.svg'
        figure.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="30" height="30"><circle cx="15" cy="15" r="10"/></svg>', encoding='utf-8')
        item = {'id': 'P1', 'statement': 'A measured difference', 'status': 'observed',
                'calculation': 'Synthetic test calculation', 'limitations': 'Synthetic only',
                'evidence': [{'path': evidence.name, 'sha256': sha256(evidence),
                              'locator': {'type': 'json_pointer', 'pointer': '/rows/0/value'}}],
                'figures': [{'path': figure.name, 'sha256': sha256(figure), 'caption': 'Synthetic figure'}]}
        return {'schema_version': 2, 'analysis_status': 'completed', 'question': 'Compare measurements',
                'scope': 'Synthetic data', 'measurement_contract': 'Declared outcome', 'phenomena': [item],
                'findings': [], 'limitations': 'No causal identification', 'reproduce': 'Run this test'}

    def finding(self, report, status='candidate'):
        finding = copy.deepcopy(report['phenomena'][0])
        finding.update(id='H1', status=status, explains=['P1'], priority=1, priority_reason='Review evidence',
                       mechanism='Candidate explanation', support='Measured association',
                       counterevidence='Alternative remains', next_test='Independent measurement')
        report['findings'].append(finding)
        return finding

    def test_workbook_audit_supports_configured_header_row(self):
        path = self.root / 'preamble.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Telemetry'
        ws.append(['Instrument export'])
        ws.append(['elapsed', 'outcome'])
        ws.append([12, 42])
        wb.save(path)
        wb.close()
        result = audit(path, header_row=2, sheet='Telemetry')
        self.assertEqual(result['sheets'][0]['rows_scanned'], 1)
        self.assertEqual(result['sheets'][0]['columns']['outcome']['max'], 42)

    def test_decimal_interval_boundary(self):
        rows = build(self.request([[.3, .4, 10, 0, 1, 'data']], width=.1))['windows']
        self.assertEqual(len(rows), 1)
        self.assertAlmostEqual(rows[0]['start_s'], .3)
        self.assertAlmostEqual(rows[0]['value'], 100)

    def test_event_at_decimal_boundary_uses_right_bin(self):
        q = self.request([[.3, None, 1, 0, 1, 'data']], event=True, kind='event_count', width=.1)
        rows = build(q)['windows']
        self.assertEqual(rows[2]['value'], 0)
        self.assertAlmostEqual(rows[3]['value'], 10)

    def test_terminal_boundary_cannot_silently_drop_event(self):
        for event, end in [(0.3, 0.1 * 3), (math.nextafter(.3, -math.inf), .3)]:
            with self.subTest(event=event, end=end):
                q = self.request([[event, None, 1, 0, 1, 'data']], event=True, kind='event_count', width=.1)
                q['sources'][0]['time']['observation_interval'] = [0, end]
                with self.assertRaises(ValueError):
                    build(q)

    def test_accepted_event_count_is_conserved_near_terminal_boundary(self):
        q = self.request([[.3, None, 1, 0, 1, 'data']], event=True, kind='event_count', width=.1)
        q['sources'][0]['time']['observation_interval'] = [0, .4]
        rows = build(q)['windows']
        self.assertEqual(sum(r['aggregation_numerator'] for r in rows), 1)
        self.assertEqual(sum(r['valid_observations'] for r in rows), 1)

    def test_reviewed_rate_contract_accepts_equivalent_row_models(self):
        a = summarize(self.windows([10]))[0]
        a.update(kind='interval_amount', measurement_id='reviewed_captured_octets_per_second')
        b = {**a, 'entity': 'B', 'run_id': 'b', 'kind': 'event_sum', 'mean': 8}
        self.assertEqual(paired_compare([a, b], 'A')[0]['mean_target_minus_reference'], -2)
        b['kind'] = 'event_mean'
        with self.assertRaises(ValueError):
            paired_compare([a, b], 'A')

    def test_full_coverage_survives_roundoff(self):
        rows = build(self.request([[.4, .5, 10, 0, 1, 'data']], width=.1))['windows']
        self.assertEqual(rows[0]['coverage'], 1)
        self.assertAlmostEqual(summarize(rows, min_coverage=1)[0]['mean'], 100)

    def test_missing_event_sum_is_unknown(self):
        for values in [[None], [10, None]]:
            with self.subTest(values=values):
                q = self.request([[.1 + i * .1, None, v, 0, 1, 'data'] for i, v in enumerate(values)], event=True, kind='event_sum')
                row = build(q)['windows'][0]
                self.assertIsNone(row['value'])
                self.assertEqual(row['invalid_observations'], 1)
                self.assertTrue(row['invalid_source_rows'])

    def test_complete_empty_capture_has_zero_counts(self):
        q = self.request([], event=True, kind='event_count')
        q['sources'][0]['time']['observation_interval'] = [0, 2]
        self.assertEqual([r['value'] for r in build(q)['windows']], [0, 0])

    def test_incomplete_empty_capture_is_missing(self):
        q = self.request([], event=True, kind='event_count', complete=False)
        row = build(q)['windows'][0]
        self.assertIsNone(row['value'])
        self.assertEqual(row['coverage'], 0)

    def test_valid_filter_with_no_events_is_zero(self):
        q = self.request([[.1, None, 3, 0, 1, 'management']], event=True, kind='event_count')
        q['sources'][0]['where'] = {'kind': ['data']}
        self.assertEqual(build(q)['windows'][0]['value'], 0)

    def test_unknown_filter_column_is_rejected_before_filter(self):
        q = self.request([[.1, None, 3, 0, 1, 'data']], event=True, kind='event_count')
        q['sources'][0]['where'] = {'misspelled': [None]}
        with self.assertRaises(ValueError):
            build(q)

    def test_minimum_trial_support(self):
        q = self.request([[0, 1, 0, 1, 2, 'data']], kind='ratio')
        q['sources'][0]['metrics'][0]['min_denominator'] = 10
        row = build(q)['windows'][0]
        self.assertIsNone(row['value'])
        self.assertEqual(row['aggregation_denominator'], 2)
        self.assertEqual(row['status'], 'insufficient_sample_support')

    def test_minimum_observations_excludes_invalid_flags(self):
        q = self.request([[.1, None, 1, 0, 1, 'data'], [.2, None, None, 0, 1, 'data']], event=True, kind='event_fraction')
        q['sources'][0]['metrics'][0]['min_observations'] = 2
        self.assertIsNone(build(q)['windows'][0]['value'])

    def test_capture_completeness_requires_boolean(self):
        q = self.request([], event=True, kind='event_count')
        q['sources'][0]['capture_complete'] = 1
        with self.assertRaises(ValueError):
            build(q)

    def test_empty_source_list_rejected(self):
        with self.assertRaises(ValueError):
            build({'schema_version': 1, 'window_s': 1, 'sources': []})

    def test_empty_workbook_has_clear_header_error(self):
        path = self.root / 'blank.xlsx'
        openpyxl.Workbook().save(path)
        with self.assertRaises(ValueError):
            list(table_rows(path))

    def test_different_measurement_weighting_cannot_be_paired(self):
        a = summarize(self.windows([10]))[0]
        b = {**a, 'entity': 'B', 'run_id': 'b', 'kind': 'event_mean'}
        with self.assertRaises(ValueError):
            paired_compare([a, b], 'A')

    def test_different_metric_definition_cannot_be_paired(self):
        a = summarize(self.windows([10]))[0]
        b = {**a, 'entity': 'B', 'run_id': 'b', 'definition': 'Sender bytes rather than delivered bytes'}
        with self.assertRaises(ValueError):
            paired_compare([a, b], 'A')

    def test_explicit_shared_measurement_identity(self):
        a = summarize(self.windows([10]))[0]
        a['measurement_id'] = 'reviewed_common_outcome'
        b = {**a, 'entity': 'B', 'run_id': 'b', 'definition': 'Different export description, same reviewed contract'}
        self.assertEqual(paired_compare([a, b], 'A')[0]['n_pairs'], 1)

    def test_unknown_identity_case_and_whitespace(self):
        for unknown in ['Unknown', ' N/A ', 'unspecified']:
            a = summarize(self.windows([10]))[0]
            a['pair_id'] = unknown
            self.assertEqual(paired_compare([a], 'A')[0]['status'], 'blocked_no_eligible_pairs')

    def test_comparison_keeps_condition_contrasts(self):
        a = summarize(self.windows([10]))[0]
        b = {**a, 'entity': 'B', 'run_id': 'b', 'mean': 8}
        result = paired_compare([a, b], 'A')[0]
        self.assertEqual(result['condition_contrasts'][0]['target_minus_reference'], -2)
        self.assertEqual(result['pair_contrasts'][0]['pair_id'], 'p')

    def test_inconsistent_normalized_value_is_rejected(self):
        rows = self.windows([10])
        rows[0].update(aggregation_numerator=2, aggregation_denominator=1, scale=1)
        with self.assertRaises(ValueError):
            validate_windows(rows)

    def test_nonfinite_window_fields_rejected(self):
        for field in ['start_s', 'end_s', 'coverage', 'value']:
            for value in [math.nan, math.inf]:
                with self.subTest(field=field, value=value):
                    rows = self.windows([10])
                    rows[0][field] = value
                    with self.assertRaises(ValueError):
                        validate_windows(rows)

    def test_json_numeric_overflow_rejected(self):
        path = self.root / 'overflow.json'
        path.write_text('{"value":1e999}')
        with self.assertRaises(ValueError):
            read_json(path)

    def test_external_negative_fraction_rejected(self):
        rows = self.windows([-.5])
        rows[0].update(kind='ratio', aggregation_numerator=-1, aggregation_denominator=2, scale=1)
        with self.assertRaises(ValueError):
            validate_windows(rows)

    def test_external_ratio_requires_primitives(self):
        rows = self.windows([.5])
        rows[0]['kind'] = 'ratio'
        with self.assertRaises(ValueError):
            validate_windows(rows)

    def test_duplicate_json_keys_rejected(self):
        path = self.root / 'bad.json'
        path.write_text('{"value":1,"value":2}')
        with self.assertRaises(ValueError):
            read_json(path)

    def test_nonfinite_json_rejected(self):
        path = self.root / 'bad.json'
        path.write_text('{"value":NaN}')
        with self.assertRaises(ValueError):
            read_json(path)

    def test_relapse_is_not_censoring(self):
        events = detect(self.windows([10, 1, 10, 1, 10, 10]), 5, 'below', 1, .8, 2)
        self.assertEqual(events[0]['recovery_status'], 'interrupted_by_relapse')
        self.assertFalse(events[0]['recovery_censored'])
        self.assertEqual(events[0]['recovery_followup_end_s'], 3)
        self.assertEqual(events[1]['recovery_status'], 'observed')
        self.assertEqual(events[1]['recovery_start_s'], 4)

    def test_partial_window_cannot_confirm_continuous_recovery(self):
        rows = self.windows([1, 10])
        rows[1]['coverage'] = .8
        event = detect(rows, 5, 'below', 1, .8, 1)[0]
        self.assertEqual(event['recovery_status'], 'gap_censored')
        self.assertIsNone(event['recovery_start_s'])

    def test_invalid_event_policies_rejected(self):
        for value in [math.nan, math.inf, -1, 0]:
            for parameter in ['minimum', 'hold']:
                with self.subTest(value=value, parameter=parameter):
                    with self.assertRaises(ValueError):
                        detect(self.windows([1, 10]), 5, 'below', value if parameter == 'minimum' else 1, .8,
                               value if parameter == 'hold' else 1)

    def test_reversed_interval_rejected_in_detection(self):
        rows = self.windows([1])
        rows[0]['end_s'] = -1
        with self.assertRaises(ValueError):
            detect(rows, 5, 'below', 1, .8, 1)

    def test_singleton_ecdf_has_visible_jump(self):
        x, y = ecdf_points([3], [1])
        self.assertEqual(list(x), [3, 3])
        self.assertEqual(list(y), [0, 1])

    def test_plot_rejects_incompatible_contracts(self):
        rows = self.windows([10])
        other = {**rows[0], 'entity': 'B', 'run_id': 'b', 'definition': 'Another layer'}
        with self.assertRaises(ValueError):
            plot(rows + [other], self.root / 'figures')

    def test_plot_manifest_cannot_overwrite_input(self):
        path = self.root / 'figures.json'
        path.write_text(json.dumps({'windows': self.windows([3])}))
        before = sha256(path)
        result = subprocess.run([sys.executable, '-B', '-X', 'utf8', str(Path(__file__).with_name('plot_windows.py')),
                                 '--windows', str(path), '--output-dir', str(self.root)], capture_output=True, text=True, encoding='utf-8')
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('Output must not overwrite', result.stderr)
        self.assertEqual(sha256(path), before)
        self.assertFalse(list(self.root.glob('*.png')))

    def test_hardlink_output_collision_rejected(self):
        source = self.root / 'source.json'
        source.write_text('{}')
        alias = self.root / 'alias.json'
        try:
            os.link(source, alias)
        except OSError as exc:
            self.skipTest(str(exc))
        with self.assertRaises(ValueError):
            ensure_output_safe(alias, [source])

    def test_phenomenon_cannot_claim_validated_mechanism(self):
        report = self.report()
        report['phenomena'][0]['status'] = 'validated_mechanism'
        with self.assertRaises(ValueError):
            validate_evidence(report, self.root)

    def test_report_cannot_overwrite_hashed_evidence(self):
        report = self.report()
        path = self.root / 'evidence.md'
        path.write_text('Original evidence\n')
        report['phenomena'][0]['evidence'] = [{'path': path.name, 'sha256': sha256(path),
                                             'locator': {'type': 'line_range', 'start': 1, 'end': 1}}]
        before = sha256(path)
        with self.assertRaises(ValueError):
            render(report, self.root, path)
        self.assertEqual(sha256(path), before)

    def test_changed_figure_rejected(self):
        report = self.report()
        (self.root / 'plot.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg"/>')
        with self.assertRaises(ValueError):
            validate_evidence(report, self.root)

    def test_nonimage_figure_rejected(self):
        report = self.report()
        path = self.root / 'plot.svg'
        path.write_text('not an image')
        report['phenomena'][0]['figures'][0]['sha256'] = sha256(path)
        with self.assertRaises((ValueError, ET.ParseError)):
            validate_evidence(report, self.root)

    def test_nonexistent_json_locator_rejected(self):
        report = self.report()
        report['phenomena'][0]['evidence'][0]['locator']['pointer'] = '/rows/999/value'
        with self.assertRaises(ValueError):
            validate_evidence(report, self.root)

    def test_json_pointer_escaping(self):
        path = self.root / 'escaped.json'
        path.write_text('{"a/b":{"~":[42]}}')
        self.assertEqual(verify_locator(path, {'type': 'json_pointer', 'pointer': '/a~1b/~0/0'}), 'verified_json_pointer')

    def test_oversized_line_locator_rejected(self):
        path = self.root / 'lines.txt'
        path.write_text('one\ntwo\n')
        with self.assertRaises(ValueError):
            verify_locator(path, {'type': 'line_range', 'start': 1, 'end': 3})

    def test_validated_result_needs_machine_locator(self):
        report = self.report()
        finding = self.finding(report, 'validated_mechanism')
        finding['validation'] = {'design': 'A design', 'independent_replication': 'Independent runs',
                                 'reviewer_assessment': 'Assumptions assessed', 'result_path': 'evidence.json'}
        finding['evidence'][0]['locator'] = {'type': 'description', 'text': 'Somewhere in the data'}
        with self.assertRaises(ValueError):
            validate_evidence(report, self.root)

    def test_validation_rationale_is_rendered(self):
        report = self.report()
        finding = self.finding(report, 'validated_mechanism')
        finding['validation'] = {'design': 'Randomized independent intervention',
                                 'independent_replication': 'Independent replication details',
                                 'reviewer_assessment': 'Explicit assumption assessment', 'result_path': 'evidence.json'}
        output = self.root / 'diagnosis.md'
        render(report, self.root, output)
        rendered = output.read_text(encoding='utf-8')
        for text in finding['validation'].values():
            self.assertIn(text, rendered)

    def test_blocked_report_does_not_require_fake_figures(self):
        report = self.report()
        report.update(analysis_status='blocked', phenomena=[], findings=[])
        output = self.root / 'blocked.md'
        render(report, self.root, output)
        self.assertIn('blocked', output.read_text(encoding='utf-8'))

    def test_finding_links_to_observed_phenomena(self):
        report = self.report()
        finding = self.finding(report)
        finding['explains'] = ['not_observed']
        with self.assertRaises(ValueError):
            validate_evidence(report, self.root)

    def test_boolean_priority_rejected(self):
        report = self.report()
        self.finding(report)['priority'] = True
        with self.assertRaises(ValueError):
            validate_evidence(report, self.root)

    def test_boolean_sequence_rejected(self):
        with self.assertRaises(ValueError):
            decode_compressed_ba(True, '0100000000000000', 64)

    def test_duplicate_initial_retry_observation_is_flagged(self):
        row = {'capture_id': 'c', 'transmitter': 'tx', 'receiver': 'rx', 'tid': 0, 'seq': 1,
               'fragment': 0, 'time_s': 0., 'retry': 0, 'source_row': 2}
        result = observed_retry_chains([row, {**row, 'source_row': 3}], 1)
        self.assertEqual(len(result['chains']), 1)
        self.assertEqual(result['chains'][0]['duplicate_time_observations'], 1)


if __name__ == '__main__':
    unittest.main(verbosity=2)
